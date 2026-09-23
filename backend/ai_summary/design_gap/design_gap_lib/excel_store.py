from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


EXCEL_HEADERS = [
    "form_id",
    "control_number",
    "company_identifier",
    "unit_id",
    "business_process",
    "financial_year",
    "control_design_status",
    "summary",
    "ai_response_json",
    "results_json",
    "usage_json",
    "generated_at",
]


def write_excel_store(
    out_path: Path,
    *,
    control_results: list[dict[str, Any]],
    generated_at: str,
) -> Path:
    """
    One row per control: form_id + AI JSON (for later DB / UI display).
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "design_gap"

    header_font = Font(bold=True)
    for col, header in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font

    for row_idx, control in enumerate(control_results, start=2):
        ai_json = control.get("ai_response_json")
        if ai_json is None:
            # Build a storeable payload even when AI was skipped (pre-check / dry-run)
            ai_json = {
                "control_design_status": control.get("control_design_status"),
                "summary": control.get("summary"),
                "results": control.get("results") or [],
                "source": "no_openrouter_call",
            }

        values = [
            control.get("form_id"),
            control.get("control_number"),
            control.get("company_identifier"),
            control.get("unit_id"),
            control.get("business_process"),
            control.get("financial_year"),
            control.get("control_design_status"),
            control.get("summary"),
            json.dumps(ai_json, ensure_ascii=False),
            json.dumps(control.get("results") or [], ensure_ascii=False),
            json.dumps(control.get("usage") or {}, ensure_ascii=False),
            generated_at,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col in (9, 10, 11):  # JSON columns
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 48
    ws.column_dimensions["I"].width = 60
    ws.column_dimensions["J"].width = 40

    wb.save(out_path)
    return out_path
